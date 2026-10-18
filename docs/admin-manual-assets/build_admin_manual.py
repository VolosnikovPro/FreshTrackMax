from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
ASSETS = ROOT / "docs" / "admin-manual-assets"
OUTPUT = ROOT / "docs" / "FreshTrack_Инструкция_для_главного_администратора.docx"

BLUE = RGBColor(37, 99, 235)
DARK = RGBColor(15, 23, 42)
MUTED = RGBColor(71, 85, 105)
GREEN = RGBColor(22, 101, 52)
RED = RGBColor(185, 28, 28)
GOLD = RGBColor(146, 64, 14)


def set_font(run, size=None, color=None, bold=None):
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold


def shade(paragraph, fill):
    props = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    props.append(shading)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def set_cell_text(cell, text, bold=False, color=DARK, size=10):
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    set_font(r, size=size, color=color, bold=bold)


def add_page_number(paragraph):
    run = paragraph.add_run("Страница ")
    set_font(run, 9, MUTED)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def add_title(doc, title, subtitle=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(title)
    set_font(r, 27, DARK, True)
    if subtitle:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(18)
        r = p.add_run(subtitle)
        set_font(r, 12, MUTED)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    set_font(r, 17 if level == 1 else 13, BLUE if level == 1 else DARK, True)
    return p


def add_text(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.15
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        set_font(r, 10.5, DARK, True)
        r = p.add_run(text[len(bold_lead):])
        set_font(r, 10.5, DARK)
    else:
        r = p.add_run(text)
        set_font(r, 10.5, DARK)
    return p


def add_step(doc, number, title, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(f"Шаг {number}. {title}")
    set_font(r, 11, DARK, True)
    add_text(doc, text)


def add_note(doc, label, text, fill="EFF6FF", color=BLUE):
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(f"{label}. ")
    set_font(r, 10.5, color, True)
    r = p.add_run(text)
    set_font(r, 10.5, DARK)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_screenshot(doc, filename, caption):
    doc.add_picture(str(ASSETS / filename), width=Inches(6.2))
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run(caption)
    set_font(r, 9, MUTED)


def add_access_table(doc):
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    table.columns[0].width = Inches(2.0)
    table.columns[1].width = Inches(4.2)
    for cell, text in zip(table.rows[0].cells, ["Кто", "Что делает"]):
        set_cell_shading(cell, "DBEAFE")
        set_cell_text(cell, text, bold=True, color=DARK)
    rows = [
        ("Главный администратор", "Видит все объекты, создаёт пользователей и объекты, загружает риск-отчёты, смотрит партии, уведомления и историю импортов."),
        ("Администратор", "Работает с несколькими объектами, но не управляет всеми пользователями."),
        ("Кладовщик", "Работает только со своим объектом. Загружает файл расхода и списывает сырьё по правилу FEFO."),
        ("УПП и УОП", "Видят только свой объект и свой рабочий функционал."),
    ]
    for role, description in rows:
        cells = table.add_row().cells
        set_cell_text(cells[0], role, bold=True)
        set_cell_text(cells[1], description)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.5)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = header.add_run("FreshTrack | Инструкция для главного администратора")
    set_font(r, 8.5, MUTED)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_page_number(footer)

    add_title(doc, "FreshTrack", "Инструкция для главного администратора.")
    add_note(doc, "Главная мысль", "Сначала выберите объект сверху, потом выполняйте любое действие. Так данные останутся в нужном объекте.", "EFF6FF", BLUE)
    add_text(doc, "Объект - это отдельное место учёта сырья. Партия - это конкретная поставка продукта с количеством и сроком годности.")
    add_heading(doc, "С чего начать", 1)
    add_step(doc, 1, "Откройте приложение", "В верхней части экрана справа должно быть ваше имя и подпись «Главный администратор». Если подпись другая, выйдите из приложения и войдите под нужной учётной записью.")
    add_step(doc, 2, "Выберите объект", "Нажмите на название объекта рядом с логотипом FreshTrack и выберите нужный объект. После этого все разделы покажут его данные.")
    add_note(doc, "Проверка", "Перед загрузкой файла всегда посмотрите на название объекта вверху. Это самый важный шаг.", "FFF7ED", GOLD)
    add_screenshot(doc, "01-dashboard.png", "Экран дашборда. Стрелка выбора объекта находится вверху рядом с «Объект N 1».")

    add_heading(doc, "Что может главный администратор", 1)
    add_access_table(doc)
    add_text(doc, "Главный администратор не должен загружать файл расхода вместо кладовщика. Это разные операции: риск-отчёт добавляет и обновляет поставки, а файл расхода уменьшает остатки.")

    add_heading(doc, "Дашборд: быстрый осмотр", 1)
    add_text(doc, "Дашборд - это главная страница. Он помогает за несколько секунд понять, всё ли хорошо.")
    add_step(doc, 1, "Посмотрите на карточки", "«Продуктов» - сколько названий сырья есть. «Партий» - сколько поставок хранится. «Срочно» - сколько партий нужно проверить прямо сейчас.")
    add_step(doc, 2, "Проверьте красные сроки", "Красная метка означает срок до 15 дней. Откройте «Партии», найдите продукт и свяжитесь с ответственным сотрудником.")
    add_step(doc, 3, "Проверьте жёлтые сроки", "Жёлтая метка означает 16-45 дней. Это ещё не срочно, но продукт нужно запланировать к использованию.")

    add_heading(doc, "Как загрузить риск-отчёт", 1)
    add_text(doc, "Риск-отчёт загружает новые продукты и партии или обновляет уже найденные записи для выбранного объекта.")
    add_step(doc, 1, "Выберите объект", "Сверху выберите объект, для которого пришёл отчёт.")
    add_step(doc, 2, "Откройте «Импорт»", "Внизу нажмите «Ещё», затем «Импорт». В поле «Тип файла» оставьте «Риск-отчёт».")
    add_step(doc, 3, "Выберите файл", "Нажмите большую область «Выбрать XLSX-файл» и укажите файл в формате .xlsx или .xlsm.")
    add_step(doc, 4, "Нажмите «Импортировать»", "Дождитесь результата. Зеленая карточка покажет, сколько строк добавлено, обновлено или пропущено.")
    add_step(doc, 5, "Сразу проверьте результат", "Откройте «Продукты» и «Партии / Сроки». Там должны быть данные именно выбранного объекта.")
    add_note(doc, "Нельзя", "Не загружайте один и тот же файл повторно. Система распознает дубликат и не даст случайно удвоить данные.", "FEF2F2", RED)
    add_screenshot(doc, "02-import.png", "Экран импорта. Вверху уже выбран объект, внизу видна история завершённых импортов.")

    add_heading(doc, "Как проверить, что сырьё появилось", 1)
    add_step(doc, 1, "Откройте «Продукты»", "Здесь отображаются названия сырья. Используйте строку поиска, если список длинный.")
    add_step(doc, 2, "Откройте «Партии»", "Здесь отображаются количество, номер УПД и срок годности каждой поставки.")
    add_step(doc, 3, "Сверьте объект", "Если данные не видны, сначала проверьте название объекта сверху. Очень часто выбран другой объект.")
    add_screenshot(doc, "03-products.png", "Продукты выбранного объекта. Кнопка «+» нужна для ручного добавления новой записи.")
    add_screenshot(doc, "04-batches.png", "Партии и сроки. Красная метка - срочно, жёлтая - нужно запланировать использование.")

    add_heading(doc, "Пользователи: кого и куда добавить", 1)
    add_text(doc, "Управление пользователями находится в разделе «Ещё» -> «Администрирование» -> «Пользователи».")
    add_step(doc, 1, "Нажмите «Новый пользователь»", "Введите логин и понятное отображаемое имя.")
    add_step(doc, 2, "Выберите роль", "Кладовщик списывает сырьё файлом. УПП работает с финансовыми данными. УОП смотрит свой объект и данные по питанию. Администратор может переключать объекты.")
    add_step(doc, 3, "Выберите объект", "Для кладовщика, УПП и УОП это их единственный объект. Они не могут менять его сами.")
    add_step(doc, 4, "Свяжите с MAX при необходимости", "Внесите MAX User ID, если пользователь будет входить через MAX. Пароль нужен только для локального входа.")
    add_step(doc, 5, "Сохраните и проверьте", "Новый пользователь появится в списке. Кнопка с карандашом редактирует, «Откл» временно отключает доступ, корзина удаляет учётную запись.")
    add_note(doc, "Осторожно", "Перед отключением или удалением убедитесь, что это не ваша учётная запись и не единственный ответственный сотрудник объекта.", "FFF7ED", GOLD)
    add_screenshot(doc, "06-admin-users.png", "Администрирование -> Пользователи. На каждой карточке видны роль и назначенный объект.")

    add_heading(doc, "Объекты", 1)
    add_step(doc, 1, "Откройте «Администрирование»", "Перейдите на вкладку «Объекты».")
    add_step(doc, 2, "Создайте объект", "Нажмите «Новый объект», заполните название и короткий уникальный код, затем сохраните.")
    add_step(doc, 3, "Назначьте сотрудников", "После создания объекта откройте вкладку «Пользователи» и укажите этот объект для нужных сотрудников.")
    add_text(doc, "Неактивный объект не выбирается для новой работы. Старые данные остаются в базе, но сотрудникам он не доступен.")

    add_heading(doc, "Уведомления MAX", 1)
    add_text(doc, "Перед отправкой выберите объект вверху. Уведомление относится только к нему.")
    add_step(doc, 1, "Откройте «Уведомления»", "Внизу нажмите «Ещё» и выберите «Уведомления».")
    add_step(doc, 2, "Выберите тип", "Срочное уведомление - просроченные партии и партии со сроком до 15 дней. Плановое - срок от 16 до 45 дней.")
    add_step(doc, 3, "Отправьте", "Нажмите нужную кнопку и дождитесь зелёного результата с количеством отправленных сообщений.")
    add_note(doc, "Плановое", "Плановая рассылка предназначена для еженедельного контроля; её базовый день - воскресенье.", "EFF6FF", BLUE)

    add_heading(doc, "История и отмена импорта", 1)
    add_text(doc, "В «Импорте» под формой есть история операций для выбранного объекта. Она помогает понять, какой файл уже был загружен.")
    add_step(doc, 1, "Найдите нужную строку", "В истории видны тип файла, номер операции и результат.")
    add_step(doc, 2, "Отменяйте только ошибочный импорт", "Нажмите «Отменить» и напишите причину. Для файла расхода отмена восстанавливает списанные остатки. Для риск-отчёта сначала убедитесь, что вы отменяете именно нужный импорт.")
    add_note(doc, "Важно", "Не удаляйте данные вручную, если можно отменить импорт. История нужна, чтобы всегда было понятно, что и когда произошло.", "FEF2F2", RED)

    add_heading(doc, "Себестоимость и питающиеся", 1)
    add_text(doc, "Главный администратор видит сводные финансовые данные. Ввод суммы накладной и количества питающихся выполняет УПП. Это разделение нужно, чтобы финансовые цифры вводил ответственный сотрудник.")
    add_text(doc, "Сейчас экран «Питающиеся» требует корректировки доступа для роли главного администратора: раздел есть в навигации, но не открывается для `superadmin`. До исправления вносить и проверять эти данные следует через учётную запись администратора или УПП.")

    add_heading(doc, "Если появилась ошибка", 1)
    add_step(doc, 1, "Файл не загружается", "Проверьте формат: разрешены .xlsx и .xlsm. Убедитесь, что файл не больше 20 МБ.")
    add_step(doc, 2, "Сырьё не видно после импорта", "Проверьте верхний селектор объекта, затем откройте «Продукты» и «Партии». Если импорт показал ошибки, исправьте файл и загрузите новую версию.")
    add_step(doc, 3, "Не удаётся выбрать объект", "Эта функция доступна только главному администратору и администратору. Для остальных ролей это правильно: они закреплены за одним объектом.")
    add_step(doc, 4, "Нет пункта «Администрирование»", "Обновите страницу. В правом верхнем углу должна быть роль «Главный администратор». Если указана другая роль, обратитесь к главному администратору.")

    add_heading(doc, "Короткий ежедневный маршрут", 1)
    add_text(doc, "1. Выберите объект.  2. Откройте дашборд.  3. Проверьте красные и жёлтые сроки.  4. При необходимости загрузите риск-отчёт.  5. Сверьте продукты и партии.  6. Проверьте историю импорта.  7. Отправьте уведомление, если ситуация требует внимания.")
    add_note(doc, "Запомните", "Главный администратор отвечает за правильный объект, правильные доступы сотрудников и понятную историю действий. Не торопитесь: сначала проверьте объект, потом нажимайте кнопку.", "ECFDF5", GREEN)

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
