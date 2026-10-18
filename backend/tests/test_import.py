"""Tests for Google-Sheets / xlsx import (risk_report and cost_sheet)."""
from openpyxl import Workbook
from datetime import date, timedelta

# Excel serial date = days since 1899-12-30
_EXCEL_EPOCH = date(1899, 12, 30)


def _to_excel_serial(d: date) -> int:
    return (d - _EXCEL_EPOCH).days


def _make_risk_xlsx(path):
    wb = Workbook()
    ws = wb.active
    ws.append(["", "Риск-Отчет"])
    ws.append(["Наименование товаров", "Ед. Изм.", "Количество", "Срок годности", "Комментарий"])
    ws.append(["Сыр Гауда", "кг", 5, _to_excel_serial(date.today() + timedelta(days=10)), "проверить"])
    ws.append(["Масло сливочное", "кг", 3, _to_excel_serial(date.today() + timedelta(days=30)), ""])
    wb.save(path)


def _make_cost_xlsx(path):
    wb = Workbook()
    ws = wb.active
    ws.append(["Склад Асача"])
    ws.append(["Себестоимость"])
    ws.append(["Сумма перемещения"])
    ws.append(["Кол-во человек"])
    ws.append(["Код", "Товар", "Ед. изм.", "Цена", "Остаток", "1", "2"])
    ws.append([""])  # separator
    ws.append(["1", "Сахар", "кг", 120, 50])
    ws.append(["2", "Мука", "кг", 80, 30])
    wb.save(path)


def test_import_risk_report_creates_products_and_batches(seeded_client, auth_header, tmp_path):
    _make_risk_xlsx(tmp_path / "risk.xlsx")
    with open(tmp_path / "risk.xlsx", "rb") as f:
        resp = seeded_client.post(
            "/api/sheets/import/file?import_type=risk_report",
            headers=auth_header,
            files={"file": ("risk.xlsx", f, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
    assert resp.status_code in (200, 201), resp.text
    data = resp.json()
    assert data["imported"] == 2
    assert data["errors"] == []

    products = seeded_client.get("/api/products/?active_only=false", headers=auth_header).json()
    names = {p["name"] for p in products}
    assert "Сыр Гауда" in names and "Масло сливочное" in names

    batches = seeded_client.get("/api/batches/?include_archived=true", headers=auth_header).json()
    assert len(batches) >= 2  # Молоко from elsewhere + 2 here


def test_import_cost_sheet_creates_costs(seeded_client, auth_header, tmp_path):
    _make_cost_xlsx(tmp_path / "cost.xlsx")
    with open(tmp_path / "cost.xlsx", "rb") as f:
        resp = seeded_client.post(
            "/api/sheets/import/file?import_type=cost_sheet",
            headers=auth_header,
            files={"file": ("cost.xlsx", f, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
    assert resp.status_code in (200, 201), resp.text
    data = resp.json()
    assert data["imported"] == 2
    assert data["errors"] == []

    costs = seeded_client.get("/api/costs/", headers=auth_header).json()
    names = {c["product_name"] for c in costs}
    assert "Сахар" in names and "Мука" in names


def test_import_cost_sheet_ignores_hidden_empty_sheet(seeded_client, auth_header, tmp_path):
    wb = Workbook()
    hidden = wb.active
    hidden.title = "Лист1"
    hidden.sheet_state = "hidden"
    ws = wb.create_sheet("Сроки ТЕСТ")
    ws.append([])
    ws.append([])
    ws.append([])
    ws.append([])
    ws.append(["Код", "Товар", "Ед. изм.", "Цена"])
    ws.append([])
    ws.append([1, "Сахар", "кг", 120])
    ws.append([2, "Мука", "кг", 80])
    path = tmp_path / "prices-hidden-first-sheet.xlsx"
    wb.save(path)

    with path.open("rb") as stream:
        response = seeded_client.post(
            "/api/sheets/import/file?import_type=cost_sheet",
            headers=auth_header,
            files={
                "file": (
                    path.name,
                    stream,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )

    assert response.status_code == 200, response.text
    assert response.json()["imported"] == 2


def test_import_invalid_type_rejected(seeded_client, auth_header, tmp_path):
    _make_risk_xlsx(tmp_path / "risk.xlsx")
    with open(tmp_path / "risk.xlsx", "rb") as f:
        resp = seeded_client.post(
            "/api/sheets/import/file?import_type=bogus",
            headers=auth_header,
            files={"file": ("risk.xlsx", f, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
    assert resp.status_code == 400


def test_upp_enters_invoice_and_headcount_and_calculates_decimal_cost(seeded_client, auth_header):
    seeded_client.post("/api/admin/users", headers=auth_header, json={"username": "upp-fin", "role": "manager", "password": "pass"})
    token = seeded_client.post("/api/auth/login", json={"username": "upp-fin", "password": "pass"}).json()["access_token"]
    auth_header = {"Authorization": f"Bearer {token}"}
    response = seeded_client.post(
        "/api/costs/financial",
        headers=auth_header,
        json={"date": "2026-08-30", "invoice_amount_net": "100.00", "headcount": 3},
    )
    assert response.status_code == 201, response.text
    assert response.json()["daily_cost"] == "33.33"

    rows = seeded_client.get("/api/costs/financial?month=2026-08", headers=auth_header)
    assert rows.status_code == 200
    assert rows.json()["average"] == 33.33


def test_consumption_file_is_atomic_fefo_and_duplicate_protected(seeded_client, auth_header, make_future_date, tmp_path):
    admin_header = auth_header
    seeded_client.post("/api/admin/users", headers=auth_header, json={"username": "klad-file", "role": "storekeeper", "password": "pass"})
    token = seeded_client.post("/api/auth/login", json={"username": "klad-file", "password": "pass"}).json()["access_token"]
    auth_header = {"Authorization": f"Bearer {token}"}
    product = seeded_client.post("/api/products/", headers=auth_header, json={"name": "Файл расход", "unit": "кг"}).json()
    seeded_client.post("/api/batches/", headers=auth_header, json={"product_id": product["id"], "quantity": 5, "unit": "кг", "expiry_date": make_future_date(2), "upd_number": "УПД-1"})
    wb = Workbook()
    ws = wb.active
    ws.append(["Дата", "Товар", "Ед. изм.", "Количество"])
    ws.append([date.today(), "Файл расход", "кг", 3])
    path = tmp_path / "consumption.xlsx"
    wb.save(path)
    with path.open("rb") as stream:
        first = seeded_client.post("/api/sheets/import/file?import_type=consumption_file", headers=auth_header, files={"file": (path.name, stream, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert first.status_code == 200, first.text
    batches = seeded_client.get("/api/batches/?include_archived=true", headers=auth_header).json()
    assert batches[0]["quantity"] == 2
    with path.open("rb") as stream:
        second = seeded_client.post("/api/sheets/import/file?import_type=consumption_file", headers=auth_header, files={"file": (path.name, stream, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert second.status_code == 409, second.text
    run = seeded_client.get("/api/sheets/import/history", headers=admin_header).json()[0]
    cancelled = seeded_client.post(f"/api/sheets/import/{run['id']}/cancel", headers=admin_header)
    assert cancelled.status_code == 200, cancelled.text
    assert seeded_client.get("/api/batches/?include_archived=true", headers=auth_header).json()[0]["quantity"] == 5
    with path.open("rb") as stream:
        replay_without_confirmation = seeded_client.post("/api/sheets/import/file?import_type=consumption_file", headers=auth_header, files={"file": (path.name, stream, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert replay_without_confirmation.status_code == 409


def test_consumption_file_ignores_blank_rows(seeded_client, auth_header, make_future_date, tmp_path):
    seeded_client.post(
        "/api/admin/users",
        headers=auth_header,
        json={"username": "klad-blank-rows", "role": "storekeeper", "password": "pass"},
    )
    token = seeded_client.post(
        "/api/auth/login", json={"username": "klad-blank-rows", "password": "pass"}
    ).json()["access_token"]
    storekeeper_header = {"Authorization": f"Bearer {token}"}
    product = seeded_client.post(
        "/api/products/",
        headers=storekeeper_header,
        json={"name": "Пустые строки тест", "unit": "кг"},
    ).json()
    seeded_client.post(
        "/api/batches/",
        headers=storekeeper_header,
        json={
            "product_id": product["id"],
            "quantity": 5,
            "unit": "кг",
            "expiry_date": make_future_date(10),
        },
    )

    wb = Workbook()
    ws = wb.active
    ws.append(["Дата", "Товар", "Ед. изм.", "Количество"])
    ws.append([date.today(), "Пустые строки тест", "кг", 2])
    ws["D6"].number_format = "0.00"  # A formatted blank tail extends the sheet range.
    ws["A6"] = "   "  # Whitespace-only rows are blank for import purposes.
    path = tmp_path / "consumption-blank-rows.xlsx"
    wb.save(path)

    with path.open("rb") as stream:
        response = seeded_client.post(
            "/api/sheets/import/file?import_type=consumption_file",
            headers=storekeeper_header,
            files={
                "file": (
                    path.name,
                    stream,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )

    assert response.status_code == 200, response.text
    assert response.json()["imported"] == 1


def test_consumption_file_matches_product_with_outer_spaces_in_database(
    seeded_client, auth_header, make_future_date, tmp_path
):
    """Import should find legacy product names with outer whitespace."""
    seeded_client.post(
        "/api/admin/users",
        headers=auth_header,
        json={"username": "klad-trim-db", "role": "storekeeper", "password": "pass"},
    )
    token = seeded_client.post(
        "/api/auth/login", json={"username": "klad-trim-db", "password": "pass"}
    ).json()["access_token"]
    header = {"Authorization": f"Bearer {token}"}

    product = seeded_client.post(
        "/api/products/", headers=header, json={"name": "  Мука с пробелами  ", "unit": "кг"}
    ).json()
    assert product["name"] == "  Мука с пробелами  "
    seeded_client.post(
        "/api/batches/",
        headers=header,
        json={
            "product_id": product["id"],
            "quantity": 5,
            "unit": "кг",
            "expiry_date": make_future_date(10),
        },
    )

    wb = Workbook()
    ws = wb.active
    ws.append(["Дата", "Товар", "Ед. изм.", "Количество"])
    ws.append([date.today(), "Мука с пробелами", "кг", 2])
    path = tmp_path / "trim-consumption.xlsx"
    wb.save(path)

    with path.open("rb") as stream:
        response = seeded_client.post(
            "/api/sheets/import/file?import_type=consumption_file",
            headers=header,
            files={"file": (path.name, stream, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )

    assert response.status_code == 200, response.text


def test_audit_xlsx_serializes_import_json(seeded_client, auth_header, tmp_path):
    """Audit export must not pass Python dictionaries directly to openpyxl."""
    _make_risk_xlsx(tmp_path / "audit-risk.xlsx")
    with (tmp_path / "audit-risk.xlsx").open("rb") as stream:
        imported = seeded_client.post(
            "/api/sheets/import/file?import_type=risk_report",
            headers=auth_header,
            files={"file": ("audit-risk.xlsx", stream, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
    assert imported.status_code == 200, imported.text

    exported = seeded_client.get("/api/admin/audit/export?format=xlsx", headers=auth_header)

    assert exported.status_code == 200
    assert exported.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


def test_import_cancellation_statement_locks_the_run():
    """Two administrators must not compensate the same import concurrently."""
    from app.api.routes.sheets import import_run_for_cancel_statement

    statement = import_run_for_cancel_statement(1)

    assert statement._for_update_arg is not None


def test_cancel_risk_report_reverses_created_batches(seeded_client, auth_header, tmp_path):
    """Cancelling a risk report must remove the stock it introduced."""
    _make_risk_xlsx(tmp_path / "cancel-risk.xlsx")
    with (tmp_path / "cancel-risk.xlsx").open("rb") as stream:
        imported = seeded_client.post(
            "/api/sheets/import/file?import_type=risk_report",
            headers=auth_header,
            files={"file": ("cancel-risk.xlsx", stream, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
    assert imported.status_code == 200, imported.text
    run = seeded_client.get("/api/sheets/import/history", headers=auth_header).json()[0]

    cancelled = seeded_client.post(f"/api/sheets/import/{run['id']}/cancel", headers=auth_header)

    assert cancelled.status_code == 200, cancelled.text
    names = {batch["product_name"] for batch in seeded_client.get("/api/batches/", headers=auth_header).json()}
    assert "Сыр Гауда" not in names
    assert "Масло сливочное" not in names
